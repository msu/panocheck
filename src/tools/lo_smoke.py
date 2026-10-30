#!/usr/bin/env python3
"""Build files with real writers, fix them, and open the output in headless LibreOffice."""
import pathlib, subprocess, sys, tempfile

ROOT = pathlib.Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src/python"))
import panocheck
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from pptx import Presentation

VALUES = {"doc.title.missing": "A title", "doc.lang.missing": "en-US", "slide.title.missing": "A title",
          "table.caption.missing": "A caption", "heading.nesting": 2}
SKIP = {"list.not_marked"}


def build(d):
    doc = Document()
    doc.add_heading("Syllabus", 1)
    doc.add_heading("Deep", 3)
    doc.add_paragraph("Body ").add_run("tiny").font.size = Pt(6)
    doc.add_table(2, 2).cell(0, 0).text = "A"
    doc.core_properties.title = ""
    doc.save(d / "real.docx")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, 100, 100).text_frame.text = "No title"
    deck.save(d / "real.pptx")
    book = Workbook()
    book.active["B2"] = "x"
    book.save(d / "real.xlsx")


with tempfile.TemporaryDirectory() as tmp:
    d = pathlib.Path(tmp)
    build(d)
    for name in ("real.docx", "real.pptx", "real.xlsx"):
        path = d / name
        before = panocheck.scan(path).findings
        fixes = [{"check_id": f.check_id, "locator": f.locator, "value": VALUES.get(f.check_id)}
                 for f in before if f.fixable and f.check_id not in SKIP]
        out = d / f"fixed-{name}"
        out.write_bytes(panocheck.fix(path, fixes))
        left = {f.check_id for f in panocheck.scan(out).findings} - SKIP
        if left:
            sys.exit(f"The fix of {name} failed because these findings remain: {sorted(left)}. Correct the fix.")
        if "--no-office" not in sys.argv: subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, str(out)], check=True, timeout=180)
        if "--no-office" not in sys.argv and not out.with_suffix(".pdf").exists():
            sys.exit(f"LibreOffice failed to open {out.name}. Examine the fixes.")
        print("ok", name, sorted(f["check_id"] for f in fixes))
